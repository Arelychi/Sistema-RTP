import json
from io import BytesIO

from openpyxl import load_workbook

import app.web as web
from app.web import app


def force_memory_mode(monkeypatch):
    def unavailable_database():
        raise RuntimeError("PostgreSQL no disponible")

    monkeypatch.setattr(web, "initialize_schema", unavailable_database)
    web.memory_records.clear()


def post_record(client, **overrides):
    payload = {
        "name": "Arely",
        "credential": "AR-01",
        "route": "Ruta Centro",
        "shift": "1",
        "nomenclature": "A",
        "month": 8,
        "year": 2026,
        "work_days": [1, 2],
        "rest_days": [],
        "daily_assignments": {"1": "A1", "2": "A1"},
    }
    payload.update(overrides)
    return client.post("/api/records", json=payload)


def test_daily_details_are_updated_without_changing_other_days(monkeypatch):
    force_memory_mode(monkeypatch)
    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        first = post_record(client)
        assert first.status_code == 201

        edited = post_record(
            client,
            work_days=[1, 2],
            daily_assignments={"1": "B2", "2": "A1"},
            daily_details={"1": {"route": "Ruta Norte", "shift": "2", "nomenclature": "B", "code": "B2"}},
        )
        assert edited.status_code == 201
        saved = edited.get_json()
        assert saved["daily_details"]["1"]["route"] == "Ruta Norte"
        assert saved["daily_details"]["1"]["shift"] == "2"
        assert saved["daily_assignments"]["2"] == "A1"


def test_deleted_day_does_not_return_from_memory_merge(monkeypatch):
    force_memory_mode(monkeypatch)
    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        post_record(client)
        deleted = post_record(client, deleted_days=[1], daily_assignments={"2": "A1"})
        assert deleted.status_code == 201
        saved = deleted.get_json()
        assert "1" not in saved["daily_assignments"]
        assert 1 not in saved["work_days"]
        assert "2" in saved["daily_assignments"]

        response = client.get("/api/records?month=8&year=2026")
        assert response.status_code == 200
        assert "1" not in response.get_json()[0]["daily_assignments"]


def test_same_credential_is_not_duplicated_in_table(monkeypatch):
    force_memory_mode(monkeypatch)
    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        first = post_record(client)
        assert first.status_code == 201

        repeated = post_record(client, name="Arely", credential="AR-01", route="Ruta Centro", shift="1", nomenclature="A")
        assert repeated.status_code == 201
        assert repeated.get_json()["updated_existing"] is True

        records = client.get("/api/records?month=8&year=2026")
        assert records.status_code == 200
        assert len(records.get_json()) == 1
        assert records.get_json()[0]["credential"] == "AR-01"


def test_switching_work_to_rest_or_vacation_replaces_previous_assignment(monkeypatch):
    force_memory_mode(monkeypatch)
    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        first = post_record(client)
        assert first.status_code == 201

        rest_update = post_record(client, work_days=[], rest_days=[1], daily_assignments={}, deleted_days=[])
        assert rest_update.status_code == 201
        rest_saved = rest_update.get_json()
        assert 1 in rest_saved["rest_days"]
        assert 1 not in rest_saved["work_days"]
        assert "1" not in rest_saved["daily_assignments"]

        vacation_update = post_record(client, work_days=[], rest_days=[], vacation_days=[1], daily_assignments={"1": "V"}, deleted_days=[])
        assert vacation_update.status_code == 201
        vacation_saved = vacation_update.get_json()
        assert 1 not in vacation_saved["work_days"]
        assert 1 not in vacation_saved["rest_days"]
        assert vacation_saved["daily_assignments"]["1"] == "V"


def test_delete_record_by_credential(monkeypatch):
    force_memory_mode(monkeypatch)
    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        created = post_record(client)
        assert created.status_code == 201

        deleted = client.post("/api/records/delete", json={"credential": "AR-01", "month": 8, "year": 2026})
        assert deleted.status_code == 200
        payload = deleted.get_json()
        assert payload["deleted"] is True

        records = client.get("/api/records?month=8&year=2026")
        assert records.status_code == 200
        assert records.get_json() == []


def test_schedule_records_follow_latest_controller_name(monkeypatch):
    force_memory_mode(monkeypatch)
    web.memory_controllers = [{"id": 5, "nombre": "Ana García", "credencial": "AR-01", "sexo": "Femenino", "mod1": 1}]
    web.memory_records = [{
        "id": 1,
        "name": "Ana Vieja",
        "credential": "AR-01",
        "shift": "1",
        "route": "Ruta Centro",
        "nomenclature": "A",
        "month": 8,
        "year": 2026,
        "work_days": [1],
        "rest_days": [],
        "daily_assignments": {"1": "A1"},
        "daily_details": {"1": {"route": "Ruta Centro", "shift": "1", "nomenclature": "A", "code": "A1"}},
        "module_number": 1,
    }]

    records = web.records_for(8, 2026, 1)

    assert records[0]["name"] == "Ana García"
    assert records[0]["credential"] == "AR-01"


def test_excel_export_has_summary_matrix_structure(monkeypatch):
    force_memory_mode(monkeypatch)
    web.memory_controllers = [{"id": 5, "nombre": "Ana García", "credencial": "AR-01", "sexo": "Femenino", "mod1": 1}]
    web.memory_records = [{
        "id": 1,
        "name": "Ana Vieja",
        "credential": "AR-01",
        "shift": "1",
        "route": "Ruta Centro",
        "nomenclature": "A",
        "month": 8,
        "year": 2026,
        "work_days": [1, 2],
        "rest_days": [],
        "daily_assignments": {"1": "A1", "2": "V"},
        "daily_details": {
            "1": {"route": "Ruta Centro", "shift": "1", "nomenclature": "A", "code": "A1"},
            "2": {"route": "Ruta Centro", "shift": "1", "nomenclature": "A", "code": "V"},
        },
        "module_number": 1,
    }]

    with app.test_client() as client:
        with client.session_transaction() as session:
            session["account"] = {"username": "user1", "module_number": 1, "is_admin": False}

        response = client.get("/export/excel?month=8&year=2026")
        assert response.status_code == 200

        workbook = load_workbook(filename=BytesIO(response.data))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]

        assert "NOM" in headers
        assert "RUTA" in headers
        assert "Origen" in headers
        assert "Destino" in headers
        assert "NOMBRE DEL CONTROLADOR" in headers
        assert "TURNO" in headers
        assert sheet.max_column >= 9
